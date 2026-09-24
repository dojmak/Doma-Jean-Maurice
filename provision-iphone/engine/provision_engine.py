#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
provision_engine.py — Moteur de calcul des provisions congés payés & retraite (IFC) CMPC
v1.2.0 — 2026-07-13
Réplique ligne à ligne les formules du classeur modèle Provision_Retraite_Congé_exo_MM AA.xlsm
(onglets IAR / C / R), CORRIGE les anomalies de données connues (code L/E vide → traité comme
Local + signalé ; salariés perdus par les sous-totaux d'impression → réintégrés), recalcule des
totaux propres, cross-check contre les valeurs en cache du fichier ([L-6]), et produit le classeur PROVISION_RC_AAAA-MM.xlsx (SYNTHESE, CONGES,
RETRAITE, OD, CONTROLES).

v1.2.0 (13/07/2026) : ajout d'un garde-fou anti-cache-vide (check_cache_health) — détecte un
classeur édité programmatiquement (openpyxl) puis jamais recalculé (cache de formules externes/
internes vidé) et bloque plutôt que de produire une provision retraite silencieusement faussée.
Voir SKILL.md § Piège critique — édition programmatique du classeur.

Usage:
  python provision_engine.py <fichier_modele.xlsm> [options]
Options:
  --out DIR              dossier de sortie (défaut: cwd)
  --retours-sheet NAME   onglet des retours congés (défaut: auto-détection via formule IAR!G12)
  --cnps-sheet NAME      onglet CNPS (défaut: auto-détection nom 'cnps*' ou A1='NUMERO CNPS')
  --prev-conges X        provision congés brute au GL avant OD (défaut 0)
  --prev-chsoc-conges X  charges sociales s/congés au GL avant OD (défaut 0)
  --prev-retraite X      provision retraite (196100) au GL avant OD (défaut 0)
  --tolerance X          tolérance cross-check par ligne en FCFA (défaut 1.0)
"""
import argparse, calendar, re, sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------- Paramètres normatifs répliqués du fichier modèle (méthode validée CMPC) ----------
PLAFOND_CR = 3_375_000      # plafond CNPS retraite (décret 2022-986) — appliqué par le modèle
TAUX_CR = 0.077             # CNPS caisse de retraite part patronale
PLAFOND_AT_PF = 75_000      # plafond AT / PF appliqué par le modèle
TAUX_PF = 0.0575            # prestations familiales
TAUX_ITS_LOCAL = 0.012      # ITS employeur personnel local (modèle)
TAUX_ITS_EXPAT = 0.104      # ITS employeur personnel expatrié (modèle)
TAUX_FDFP = 0.016           # FDFP 1,2% + 0,4% (modèle)

# Comptes OD — grounding PK pcsyscohada-classe1/classe4 + referentiel_comptes_cmpc.csv (r4+)
CPT = {
    "conges_charge":      ("661300", "Congés payés", "CONFIRME_D365"),
    "conges_dette":       ("4281xx", "Dettes provisionnées pour congés à payer", "A_DETERMINER_D365"),
    "chsoc_charge":       ("664100", "Charges sociales sur rémunération", "CONFIRME_D365"),
    "chsoc_dette":        ("4382xx", "Charges sociales sur congés à payer", "A_DETERMINER_D365"),
    "retraite_dotation":  ("691100", "Dotations provisions risques et charges", "CONFIRME_D365"),
    "retraite_provision": ("196100", "Provision pour pensions", "CONFIRME_D365"),
    "retraite_reprise":   ("791100", "Reprise de provision pour risques et charges", "CONFIRME_D365"),
}

FMT_NUM = "# ##0"
FMT_DEC = "# ##0,00"
OLIVE = "3D4A1F"; CREAM = "F0EDE4"

# ------------------------------- fonctions Excel répliquées -------------------------------
def days360_us(d1, d2):
    """DAYS360 méthode US/NASD (défaut Excel)."""
    sd, sm, sy = d1.day, d1.month, d1.year
    ed, em, ey = d2.day, d2.month, d2.year
    if sd == calendar.monthrange(sy, sm)[1]:
        sd = 30
    if ed == 31:
        if sd < 30:
            em += 1
            if em == 13:
                em = 1; ey += 1
            ed = 1
        else:
            ed = 30
    return (ey - sy) * 360 + (em - sm) * 30 + (ed - sd)

def to_date(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    return None

def mois_correspondants(entree, arrete):
    """IAR!E : mois de salaire de l'exercice en cours."""
    base = date(arrete.year, 1, 1)
    if entree and entree > base:
        base = entree
    return max(0, (arrete.year - base.year) * 12 + (arrete.month - base.month) + 1)

def jours_conges_local(code, anc, mois_retour):
    """C!I : jours acquis local (2,2 j × 1,25 par mois + bonus ancienneté)."""
    if code == "E": return 0.0
    if anc >= 25: b = 7/12
    elif anc >= 20: b = 5/12
    elif anc >= 15: b = 3/12
    elif anc >= 10: b = 2/12
    elif anc >= 5:  b = 1/12
    else: return mois_retour * (1.25 * 2.2)
    return mois_retour * 1.25 * (2.2 + b)

def jours_conges_expat(code, anc, mois_retour):
    if code == "L": return 0.0
    return mois_retour * (6 if anc >= 2 else 5)

def droits_retraite(anc, sal_moyen):
    """R!L : droits simples (barème 30/35/40 % du salaire moyen mensuel par an d'ancienneté)."""
    if anc >= 10: return sal_moyen * (0.40 * anc - 0.75)
    if anc >= 5:  return sal_moyen * (0.35 * anc - 0.25)
    return anc * 0.30 * sal_moyen

def charges(l, code, taux_at):
    cr = min(l, PLAFOND_CR) * TAUX_CR
    at = min(l, PLAFOND_AT_PF) * taux_at
    pf = min(l, PLAFOND_AT_PF) * TAUX_PF
    its = l * (TAUX_ITS_LOCAL if code == "L" else TAUX_ITS_EXPAT)
    fdfp = l * TAUX_FDFP
    return cr, at, pf, its, fdfp

# ------------------------------------- lecture entrée -------------------------------------
def detect_sheets(wbf, args):
    retours = args.retours_sheet; cnps = args.cnps_sheet
    if not retours:
        f = wbf["IAR"]["G12"].value
        if isinstance(f, str):
            m = re.search(r"(?:'([^']+)'|([A-Za-z0-9_]+))!\$?G", f)
            if m: retours = m.group(1) or m.group(2)
    if not cnps:
        for ws in wbf.worksheets:
            if ws.title.lower().startswith("cnps") or ws["A1"].value == "NUMERO CNPS":
                cnps = ws.title; break
    return retours, cnps

def build_lookup(ws, key_col, val_col, row_start, row_end):
    d = {}
    for r in range(row_start, row_end + 1):
        k = ws.cell(row=r, column=key_col).value
        if isinstance(k, str) and k.strip():
            v = to_date(ws.cell(row=r, column=val_col).value)
            if k.strip() not in d and v:
                d[k.strip()] = v
    return d

def load_input(path, args):
    wbf = openpyxl.load_workbook(path, data_only=False)
    wbv = openpyxl.load_workbook(path, data_only=True)
    retours_name, cnps_name = detect_sheets(wbf, args)
    if "IAR" not in wbv.sheetnames:
        sys.exit("ERREUR: onglet IAR introuvable.")
    iar_v, iar_f = wbv["IAR"], wbf["IAR"]
    arrete = to_date(iar_v["G2"].value)
    taux_at = iar_v["G4"].value or 0.0
    if not arrete:
        sys.exit("ERREUR: date d'arrêté IAR!G2 absente.")
    lk_entree = build_lookup(wbv[cnps_name], 4, 6, 2, wbv[cnps_name].max_row) if cnps_name and cnps_name in wbv.sheetnames else {}
    lk_retour = build_lookup(wbv[retours_name], 2, 7, 14, wbv[retours_name].max_row) if retours_name and retours_name in wbv.sheetnames else {}
    emp, controles = [], []
    for r in range(12, iar_v.max_row + 1):
        nom = iar_v.cell(row=r, column=2).value
        if not (isinstance(nom, str) and nom.strip()): continue
        nom = nom.strip()
        code_raw = iar_v.cell(row=r, column=3).value
        code = str(code_raw).strip().upper() if code_raw is not None else ""
        if code not in ("L", "E"):
            controles.append((nom, f"Code L/E absent ou invalide ({code_raw!r}) : CORRIGÉ en 'L' (Local) par le moteur — dans Excel ce cas CUMULE jours local+expatrié (provision fausse) ; corriger IAR colonne C"))
            code = "L"
        sal = iar_v.cell(row=r, column=4).value
        if sal in (None, "", 0):
            controles.append((nom, "Salaire annuel absent ou nul — salarié exclu de l'effectif (règle IAR!L)"))
            continue
        # Entrée / retour : valeur littérale saisie (override) sinon lookup répliqué
        f_raw = iar_f.cell(row=r, column=6).value
        entree = to_date(f_raw) if not (isinstance(f_raw, str) and f_raw.startswith("=")) else lk_entree.get(nom)
        g_raw = iar_f.cell(row=r, column=7).value
        retour = to_date(g_raw) if not (isinstance(g_raw, str) and g_raw.startswith("=")) else lk_retour.get(nom)
        if entree is None: controles.append((nom, "Date d'entrée introuvable (lookup CNPS) — ancienneté=0, retraite=0"))
        if retour is None: controles.append((nom, "Retour congés introuvable — mois acquis congés=0"))
        # cache pour cross-check
        cache = {"E": iar_v.cell(row=r, column=5).value, "I": iar_v.cell(row=r, column=9).value, "row": r}
        emp.append(dict(nom=nom, code=code, sal=float(sal), entree=entree, retour=retour, cache=cache))
    return dict(emp=emp, arrete=arrete, taux_at=taux_at, controles=controles,
                retours=retours_name, cnps=cnps_name, wbv=wbv)

# ------------------------------- garde-fou anti-cache-vide (v1.2) -------------------------------
def check_cache_health(data):
    """Détecte un classeur édité programmatiquement (openpyxl) puis jamais recalculé : le lien
    externe cnps* et/ou les caches de formules internes sont vidés, ce qui fait perdre le lookup
    d'ancienneté pour la quasi-totalité des salariés (ancienneté->0, retraite->0 en masse) sans
    aucun message d'erreur visible. Cf. incident du 13/07/2026 (20 098 001 -> 1 850 298, -90 %)."""
    n_emp = len(data["emp"])
    if n_emp < 30:
        return  # échantillon trop petit pour un diagnostic fiable
    n_entree_absente = sum(1 for _, msg in data["controles"] if msg.startswith("Date d'entrée introuvable"))
    ratio = n_entree_absente / n_emp
    if ratio > 0.15:
        sys.exit(
            f"ERREUR (anti-cache-vide v1.2) : {n_entree_absente}/{n_emp} salariés ({ratio:.0%}) "
            "sans date d'entrée résolue -- signe probable d'un classeur édité programmatiquement "
            "(openpyxl) puis NON recalculé : le lien externe cnps* / les caches de formules sont "
            "vides. Corriger : figer les cellules à lien externe sur leurs valeurs en cache "
            "d'origine, sauvegarder, puis recalculer intégralement via LibreOffice headless "
            "(soffice --headless --convert-to xlsx:\"Calc MS Excel 2007 XML\" --outdir <dir> "
            "<fichier>) AVANT de relancer ce moteur. Voir SKILL.md § Piège critique -- édition "
            "programmatique du classeur."
        )

# --------------------------------------- calcul ---------------------------------------
def compute(data):
    arrete, taux_at = data["arrete"], data["taux_at"]
    rows = []
    for e in data["emp"]:
        mois = mois_correspondants(e["entree"], arrete)
        sal_m = e["sal"] / mois if mois else 0.0
        anc = days360_us(e["entree"], arrete) / 360 if e["entree"] else 0.0
        mois_ret = days360_us(e["retour"], arrete) / 30 if e["retour"] else 0.0
        jl = jours_conges_local(e["code"], anc, mois_ret)
        je = jours_conges_expat(e["code"], anc, mois_ret)
        prov_c = sal_m / 30 * (jl + je)
        prov_r = droits_retraite(anc, sal_m)
        ch_c = charges(prov_c, e["code"], taux_at)
        ch_r = charges(prov_r, e["code"], taux_at)
        rows.append(dict(e=e, mois=mois, sal_m=sal_m, anc=anc, mois_ret=mois_ret,
                         jl=jl, je=je, jc=jl + je, prov_c=prov_c, prov_r=prov_r,
                         ch_c=ch_c, ch_r=ch_r))
    return rows

def crosscheck(data, rows, tol):
    """[L-6] Cross-check indépendant : recalcul vs valeurs en cache du fichier, apparié par NOM.
    Détecte aussi les salariés absents des onglets C/R (décalages de références du modèle)."""
    ecarts, absents = [], []
    wbv = data["wbv"]
    caches = {}
    for sheet, lbl in (("C", "Congés"), ("R", "Retraite")):
        d = {}
        if sheet in wbv.sheetnames:
            ws = wbv[sheet]
            for r in range(7, ws.max_row + 1):
                n = ws.cell(row=r, column=1).value
                v = ws.cell(row=r, column=12).value
                if isinstance(n, str) and n.strip() and isinstance(v, (int, float)) and n.strip() not in d:
                    d[n.strip()] = float(v)
        caches[lbl] = d
    for row in rows:
        nom = row["e"]["nom"]
        for key, lbl in (("prov_c", "Congés"), ("prov_r", "Retraite")):
            d = caches[lbl]
            if not d: continue
            if nom not in d:
                absents.append((nom, lbl, row[key]))
            elif abs(row[key] - d[nom]) > tol:
                ecarts.append((nom, lbl, d[nom], row[key], row[key] - d[nom]))
    return ecarts, absents

# --------------------------------------- sortie ---------------------------------------
def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=OLIVE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def _sheet_detail(wb, name, rows, key, arrete):
    ws = wb.create_sheet(name)
    hdr = ["Noms et Prénoms", "Code", "Salaire annuel", "Mois", "Salaire moyen mensuel",
           "Entrée", "Retour congés", "Ancienneté (ans)", "Jours local", "Jours expatrié",
           "Jours total", ("Congés à payer" if key == "prov_c" else "Droits simples"),
           "CNPS CR 7,7%", "AT", "PF 5,75%", "ITS empl.", "FDFP", "Total charges"]
    ws.cell(row=1, column=1, value=f"PROVISION {'CONGES PAYES' if key=='prov_c' else 'RETRAITE (IFC)'} AU {arrete.strftime('%d/%m/%Y')}").font = Font(bold=True, size=12, color=OLIVE)
    for c, h in enumerate(hdr, 1): ws.cell(row=3, column=c, value=h)
    _style_header(ws, 3, len(hdr))
    r = 4
    tot = [0.0] * 7
    for row in rows:
        e = row["e"]; ch = row["ch_c"] if key == "prov_c" else row["ch_r"]
        prov = row[key]
        vals = [e["nom"], e["code"], e["sal"], row["mois"], row["sal_m"],
                e["entree"], e["retour"], round(row["anc"], 2), row["jl"], row["je"],
                row["jc"], prov, *ch, sum(ch)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in (3, 5, 12, 13, 14, 15, 16, 17, 18): cell.number_format = FMT_NUM
            if c in (8, 9, 10, 11): cell.number_format = FMT_DEC
            if c in (6, 7): cell.number_format = "DD/MM/YYYY"
        tot[0] += prov
        for i in range(5): tot[i + 1] += ch[i]
        tot[6] += sum(ch)
        r += 1
    ws.cell(row=r, column=1, value=f"TOTAL ({len(rows)} salariés)").font = Font(bold=True)
    for i, c in enumerate([12, 13, 14, 15, 16, 17, 18]):
        cell = ws.cell(row=r, column=c, value=tot[i])
        cell.font = Font(bold=True); cell.number_format = FMT_NUM
        cell.fill = PatternFill("solid", fgColor=CREAM)
    ws.freeze_panes = "A4"
    ws.column_dimensions["A"].width = 32
    for col in "CDEFGHIJKLMNOPQR": ws.column_dimensions[col].width = 13
    return tot

def write_output(path_out, data, rows, ecarts, absents, args):
    arrete = data["arrete"]
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    tot_c = _sheet_detail(wb, "CONGES", rows, "prov_c", arrete)
    tot_r = _sheet_detail(wb, "RETRAITE", rows, "prov_r", arrete)

    # SYNTHESE
    ws = wb.create_sheet("SYNTHESE", 0)
    ws["A1"] = f"PROVISIONS CONGES PAYES & RETRAITE — CMPC — arrêté au {arrete.strftime('%d/%m/%Y')}"
    ws["A1"].font = Font(bold=True, size=13, color=OLIVE)
    ws["A2"] = f"Effectif traité : {len(rows)} salariés · Taux AT : {data['taux_at']:.1%} · Onglet retours : {data['retours']} · Onglet CNPS : {data['cnps']}"
    lines = [
        ("", ""), ("PROVISION", "MONTANT (FCFA)"),
        ("Congés à payer (brut)", tot_c[0]),
        ("Charges sociales s/ congés (CR+AT+PF)", tot_c[1] + tot_c[2] + tot_c[3]),
        ("Charges fiscales s/ congés (ITS+FDFP)", tot_c[4] + tot_c[5]),
        ("TOTAL CONGES", tot_c[0] + tot_c[6]),
        ("Retraite — droits simples (IFC)", tot_r[0]),
        ("Charges sociales s/ retraite (CR+AT+PF)", tot_r[1] + tot_r[2] + tot_r[3]),
        ("Charges fiscales s/ retraite (ITS+FDFP)", tot_r[4] + tot_r[5]),
        ("TOTAL RETRAITE", tot_r[0] + tot_r[6]),
    ]
    for i, (a, b) in enumerate(lines, 4):
        ws.cell(row=i, column=1, value=a)
        cell = ws.cell(row=i, column=2, value=b if b != "" else None)
        if isinstance(b, float): cell.number_format = FMT_NUM
        if a.startswith("TOTAL") or a == "PROVISION":
            ws.cell(row=i, column=1).font = Font(bold=True)
            cell.font = Font(bold=True)
            ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=CREAM)
            cell.fill = PatternFill("solid", fgColor=CREAM)
    nb = len(ecarts) + len(absents)
    ws.cell(row=15, column=1, value=f"Cross-check [L-6] vs fichier modèle : {'OK — aucun écart > tolérance' if nb == 0 else str(len(ecarts)) + ' écart(s), ' + str(len(absents)) + ' salarié(s) absent(s) des onglets du fichier — voir CONTROLES'}").font = Font(italic=True, color=OLIVE if nb == 0 else "AA0000")
    ws.column_dimensions["A"].width = 45; ws.column_dimensions["B"].width = 20

    # OD
    ws = wb.create_sheet("OD")
    ws["A1"] = f"OD PROVISIONS AU {arrete.strftime('%d/%m/%Y')} — variation vs soldes GL fournis"
    ws["A1"].font = Font(bold=True, size=12, color=OLIVE)
    ws["A2"] = "⚠ À valider avant saisie D365. Comptes marqués [COMPTE À DÉTERMINER] : sous-compte D365 à confirmer (JM/DAF)."
    ws["A2"].font = Font(italic=True, color="AA0000")
    hdr = ["Compte", "Intitulé", "Débit", "Crédit", "Statut compte"]
    for c, h in enumerate(hdr, 1): ws.cell(row=4, column=c, value=h)
    _style_header(ws, 4, len(hdr))
    chsoc_c = tot_c[1] + tot_c[2] + tot_c[3]
    d_conges = tot_c[0] - args.prev_conges
    d_chsoc = chsoc_c - args.prev_chsoc_conges
    d_retraite = tot_r[0] - args.prev_retraite
    od = []
    def _line(key, deb, cred):
        num, lib, st = CPT[key]
        od.append([num, lib, deb, cred, "OK D365" if st == "CONFIRME_D365" else "[COMPTE À DÉTERMINER]"])
    if d_conges >= 0: _line("conges_charge", d_conges, 0); _line("conges_dette", 0, d_conges)
    else: _line("conges_dette", -d_conges, 0); _line("conges_charge", 0, -d_conges)
    if d_chsoc >= 0: _line("chsoc_charge", d_chsoc, 0); _line("chsoc_dette", 0, d_chsoc)
    else: _line("chsoc_dette", -d_chsoc, 0); _line("chsoc_charge", 0, -d_chsoc)
    if d_retraite >= 0: _line("retraite_dotation", d_retraite, 0); _line("retraite_provision", 0, d_retraite)
    else: _line("retraite_provision", -d_retraite, 0); _line("retraite_reprise", 0, -d_retraite)
    r = 5
    for line in od:
        for c, v in enumerate(line, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in (3, 4): cell.number_format = FMT_NUM
        r += 1
    td, tc = sum(x[2] for x in od), sum(x[3] for x in od)
    ws.cell(row=r, column=2, value="TOTAL").font = Font(bold=True)
    ws.cell(row=r, column=3, value=td).number_format = FMT_NUM
    ws.cell(row=r, column=4, value=tc).number_format = FMT_NUM
    assert abs(td - tc) < 0.01, "ARRÊT-5 : OD déséquilibrée"
    ws.cell(row=r + 2, column=1, value=f"Bouclage D=C : {td:,.2f} = {tc:,.2f} ✓ (ARRÊT-5/I3). Charges fiscales s/congés ({tot_c[4]+tot_c[5]:,.0f}) et charges s/retraite ({tot_r[6]:,.0f}) NON incluses dans l'OD — schéma PK non tranché, arbitrage requis.")
    ws.cell(row=r + 2, column=1).font = Font(italic=True, size=9)
    ws.column_dimensions["A"].width = 12; ws.column_dimensions["B"].width = 42
    for col in "CDE": ws.column_dimensions[col].width = 18

    # CONTROLES
    ws = wb.create_sheet("CONTROLES")
    ws["A1"] = "CONTRÔLES & ANOMALIES"; ws["A1"].font = Font(bold=True, size=12, color=OLIVE)
    r = 3
    ws.cell(row=r, column=1, value="Anomalies d'entrée (données manquantes) :").font = Font(bold=True); r += 1
    if not data["controles"]: ws.cell(row=r, column=1, value="aucune"); r += 1
    for nom, msg in data["controles"]:
        ws.cell(row=r, column=1, value=nom); ws.cell(row=r, column=2, value=msg); r += 1
    r += 1
    ws.cell(row=r, column=1, value=f"Cross-check [L-6] recalcul vs cache fichier (tolérance ±{args.tolerance} FCFA/ligne) :").font = Font(bold=True); r += 1
    if not ecarts: ws.cell(row=r, column=1, value="AUCUN ÉCART — recalcul conforme au fichier modèle"); r += 1
    else:
        for c, h in enumerate(["Salarié", "Provision", "Valeur fichier", "Recalcul", "Écart"], 1):
            ws.cell(row=r, column=c, value=h).font = Font(bold=True)
        r += 1
        for nom, lbl, cached, calc, d in ecarts:
            for c, v in enumerate([nom, lbl, cached, calc, d], 1):
                cell = ws.cell(row=r, column=c, value=v)
                if c >= 3: cell.number_format = FMT_DEC
            r += 1
    ws.column_dimensions["A"].width = 35; ws.column_dimensions["B"].width = 60
    for col in "CDE": ws.column_dimensions[col].width = 18

    wb.save(path_out)
    return tot_c, tot_r

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("fichier")
    ap.add_argument("--out", default=".")
    ap.add_argument("--retours-sheet", default=None)
    ap.add_argument("--cnps-sheet", default=None)
    ap.add_argument("--prev-conges", type=float, default=0.0)
    ap.add_argument("--prev-chsoc-conges", type=float, default=0.0)
    ap.add_argument("--prev-retraite", type=float, default=0.0)
    ap.add_argument("--tolerance", type=float, default=1.0)
    args = ap.parse_args()

    data = load_input(args.fichier, args)
    check_cache_health(data)
    rows = compute(data)
    ecarts, absents = crosscheck(data, rows, args.tolerance)
    out = Path(args.out) / f"PROVISION_RC_{data['arrete'].strftime('%Y-%m')}.xlsx"
    tot_c, tot_r = write_output(out, data, rows, ecarts, absents, args)
    print(f"Arrêté             : {data['arrete'].strftime('%d/%m/%Y')}")
    print(f"Effectif traité    : {len(rows)}")
    print(f"Provision congés   : {tot_c[0]:,.2f} (charges {tot_c[6]:,.2f})")
    print(f"Provision retraite : {tot_r[0]:,.2f} (charges {tot_r[6]:,.2f})")
    print(f"Cross-check [L-6]  : {'OK' if not (ecarts or absents) else f'{len(ecarts)} écart(s), {len(absents)} absent(s) des onglets du fichier — VOIR CONTROLES'}")
    print(f"Fichier produit    : {out}")
    for e in ecarts[:10]: print("  ECART:", e)
    for a in absents[:10]: print("  ABSENT DU FICHIER:", a)

if __name__ == "__main__":
    main()
