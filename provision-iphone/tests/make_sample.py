#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_sample.py — génère un classeur modèle SYNTHÉTIQUE (noms et montants fictifs) reproduisant la
structure attendue par provision_engine.py : IAR / cnps 0926 / sept / C / R.
Sert uniquement à tester la web app iPhone contre l'exécution native du moteur.

Usage : python make_sample.py <sortie.xlsx>
"""
import sys
from datetime import date

import openpyxl

ARRETE = date(2026, 9, 30)
TAUX_AT = 0.02

# nom, code L/E, salaire annuel cumulé, date d'embauche, date de retour congés
EMP = [
    ("SALARIE TEST 01", "L", 2_700_000, date(2001, 3, 15), date(2026, 2, 1)),
    ("SALARIE TEST 02", "L", 4_950_000, date(2012, 7, 1), date(2025, 12, 20)),
    ("SALARIE TEST 03", "E", 36_000_000, date(2023, 1, 9), date(2026, 6, 30)),
    ("SALARIE TEST 04", "L", 1_620_000, date(2026, 4, 1), None),
    ("SALARIE TEST 05", None, 3_150_000, date(2018, 11, 5), date(2026, 5, 31)),
    ("SALARIE TEST 06", "e", 27_000_000, date(2025, 2, 17), date(2026, 1, 31)),
    ("SALARIE TEST 07", "L", 0, date(2020, 1, 1), date(2026, 3, 1)),
    ("SALARIE TEST 08", "L", 8_100_000, date(1998, 9, 1), date(2026, 8, 31)),
]


def main(out):
    wb = openpyxl.Workbook()
    iar = wb.active
    iar.title = "IAR"
    iar["G2"] = ARRETE
    iar["G4"] = TAUX_AT
    for i, (nom, code, sal, _, _) in enumerate(EMP):
        r = 12 + i
        iar.cell(row=r, column=2, value=nom)
        iar.cell(row=r, column=3, value=code)
        iar.cell(row=r, column=4, value=sal)
        # F et G en formules de lookup, comme le modèle : le moteur réplique le lookup
        iar.cell(row=r, column=6, value=f"=VLOOKUP(B{r},'cnps 0926'!D:F,3,0)")
        iar.cell(row=r, column=7, value=f"=INDEX(sept!$G:$G,MATCH(B{r},sept!$B:$B,0))")

    cnps = wb.create_sheet("cnps 0926")
    cnps["A1"] = "NUMERO CNPS"
    for i, (nom, _, _, entree, _) in enumerate(EMP):
        cnps.cell(row=2 + i, column=4, value=nom)
        cnps.cell(row=2 + i, column=6, value=entree)

    ret = wb.create_sheet("sept")
    for i, (nom, _, _, _, retour) in enumerate(EMP):
        ret.cell(row=14 + i, column=2, value=nom)
        ret.cell(row=14 + i, column=7, value=retour)

    wb.create_sheet("C")
    wb.create_sheet("R")
    wb.save(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "sample_modele.xlsx")
